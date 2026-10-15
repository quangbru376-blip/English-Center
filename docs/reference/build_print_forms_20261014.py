from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT,WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).resolve().parents[2]
out=root/'docs/BIEU_MAU_BAN_IN_CHUYEN_NGHIEP_20261014.docx'
d=Document();s=d.sections[0]
s.page_width=Cm(21);s.page_height=Cm(29.7)
s.top_margin=Cm(1.65);s.bottom_margin=Cm(1.65);s.left_margin=Cm(2);s.right_margin=Cm(1.7)
s.header_distance=Cm(.7);s.footer_distance=Cm(.7)
s.different_first_page_header_footer=False;d.settings.odd_and_even_pages_header_footer=False
for b in list(d.styles.element.xpath('.//w:pBdr')):b.getparent().remove(b)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=d.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
 st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Times New Roman')
 for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
  st._element.get_or_add_rPr().rFonts.attrib.pop(qn('w:'+attr),None)
d.styles['Normal'].font.size=Pt(12)
d.styles['Normal'].paragraph_format.space_after=Pt(7)
d.styles['Normal'].paragraph_format.line_spacing=1.05
d.styles['Title'].font.size=Pt(17);d.styles['Title'].font.bold=True
d.styles['Title'].paragraph_format.space_after=Pt(8)
d.styles['Heading 2'].font.size=Pt(12);d.styles['Heading 2'].font.bold=True
d.styles['Heading 2'].paragraph_format.space_before=Pt(10);d.styles['Heading 2'].paragraph_format.space_after=Pt(6)
d.core_properties.title='Bộ biểu mẫu nghiệp vụ dùng cho bản in'
d.core_properties.subject='Tám biểu mẫu A4 và hướng dẫn sử dụng'

def p(text='',bold=False,size=12,align=None):
 a=d.add_paragraph();r=a.add_run(text);r.bold=bold;r.font.size=Pt(size)
 if align is not None:a.alignment=align
 return a

def small(text):return p(text,size=10)
def h(text):d.add_heading(text,2)
def line(label,n=80):
 a=p();a.add_run(label+' ').bold=True;a.add_run('.' * n);return a
def pair(left,right):
 a=p();a.paragraph_format.tab_stops.add_tab_stop(Cm(8.8))
 a.add_run(left+' ').bold=True;a.add_run('.'*25+'\t');a.add_run(right+' ').bold=True;a.add_run('.'*25)
 return a

def start(title,code,subtitle):
 if d.paragraphs:d.add_page_break()
 a=p('TRUNG TÂM TIẾNG ANH',True,11);a.paragraph_format.tab_stops.add_tab_stop(Cm(13.5));a.add_run('\t'+code).font.size=Pt(10)
 a=d.add_paragraph(title,style='Title');a.alignment=WD_ALIGN_PARAGRAPH.CENTER
 a=p(subtitle,size=11,align=WD_ALIGN_PARAGRAPH.CENTER);a.paragraph_format.space_after=Pt(15)

def table(head,rows,widths,height=.9,size=11):
 t=d.add_table(rows=1,cols=len(head));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for col,w in zip(t.columns,widths):col.width=Cm(w)
 # Explicit cell borders avoid inherited table style decoration.
 for content in [head]+rows:
  row=t.rows[0] if content is head else t.add_row()
  if content is not head:row.height=Cm(height);row.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
  pr=row._tr.get_or_add_trPr();e=OxmlElement('w:cantSplit');pr.append(e)
  if content is head:e=OxmlElement('w:tblHeader');pr.append(e)
  for j,(c,txt) in enumerate(zip(row.cells,content)):
   c.width=Cm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   c.text=str(txt);cp=c._tc.get_or_add_tcPr()
   b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');b.append(x)
   cp.append(b)
   mar=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    x=OxmlElement('w:'+edge);x.set(qn('w:w'),'80');x.set(qn('w:type'),'dxa');mar.append(x)
   cp.append(mar)
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEEEEE' if content is head else 'FFFFFF');cp.append(sh)
   for a in c.paragraphs:
    a.paragraph_format.space_before=Pt(0);a.paragraph_format.space_after=Pt(0);a.paragraph_format.line_spacing=Pt(size+1)
    no_grid=OxmlElement('w:snapToGrid');no_grid.set(qn('w:val'),'0');a._p.get_or_add_pPr().append(no_grid)
    a.alignment=WD_ALIGN_PARAGRAPH.CENTER if j==0 or content is head else WD_ALIGN_PARAGRAPH.LEFT
    for r in a.runs:r.font.size=Pt(size);r.bold=content is head
 a=d.add_paragraph();a.paragraph_format.space_after=Pt(0);a.paragraph_format.space_before=Pt(0);a.paragraph_format.line_spacing=1;a.add_run().font.size=Pt(3)
 return t

def sign(left,right=None):
 a=p('Ngày ........ tháng ........ năm ........',size=11,align=WD_ALIGN_PARAGRAPH.RIGHT)
 t=d.add_table(rows=1,cols=2 if right else 1);t.autofit=False
 for c,label in zip(t.rows[0].cells,[left,right] if right else [left]):
  c.width=Cm(8.65 if right else 17.3)
  a=c.paragraphs[0];a.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=a.add_run(label);r.bold=True
  a=c.add_paragraph('(Ký và ghi rõ họ tên)');a.alignment=WD_ALIGN_PARAGRAPH.CENTER
  for r in a.runs:r.italic=True;r.font.size=Pt(10)

start('PHIẾU ĐĂNG KÝ KHÓA HỌC','BM-TS-01','Thông tin một hồ sơ đăng ký của học viên')
pair('Mã hồ sơ','Ngày đăng ký')
h('I THÔNG TIN HỌC VIÊN')
line('Họ và tên',85);pair('Mã học viên','Ngày sinh');pair('Số điện thoại','Email')
line('Địa chỉ nếu có',83)
h('II KHÓA HỌC ĐĂNG KÝ')
pair('Mã khóa học','Tên khóa học');line('Học phí niêm yết tham khảo VND',55)
h('III THÔNG TIN THEO DÕI HỒ SƠ')
pair('Trạng thái hồ sơ','Mức trình độ')
table(['Lượt thi','Ngày thi','Điểm trên 100','Ghi chú tình trạng'],[['1','','',''],['2 nếu có','','','']],[2,3.5,3.3,8.5],.8)
pair('Điểm tốt nhất','Lớp đã xếp nếu có');line('Số hóa đơn nếu đã phát hành',58)
line('Lý do từ chối hoặc hủy nếu có',58)
small('Mỗi hồ sơ đăng ký một khóa học. Học phí chính thức được ghi trên hóa đơn phát hành khi hồ sơ được duyệt. Thông tin theo dõi được đối chiếu với hồ sơ trong hệ thống.')
sign('HỌC VIÊN')

start('DANH SÁCH LỚP','BM-ĐT-01','Danh sách học viên thuộc lớp được phân công hoặc quản lý')
pair('Mã lớp','Khóa học');pair('Chương trình','Mức trình độ')
line('Giáo viên phụ trách',75);pair('Sĩ số','Sức chứa')
table(['STT','Mã học viên','Họ và tên','Email','Số điện thoại'],[[str(i),'','','',''] for i in range(1,16)],[1.2,2,5.2,5.6,3.3],.70,10.5)
small('Thông tin liên hệ của học viên được lấy từ hồ sơ hiện tại. Danh sách được sử dụng trong phạm vi lớp phụ trách hoặc quản lý.')
sign('NGƯỜI LẬP DANH SÁCH')

start('THỜI KHÓA BIỂU CÁ NHÂN','BM-ĐT-02','Lịch học và thi cuối khóa của học viên')
line('Họ và tên học viên',78);pair('Mã học viên','Lớp hoặc khóa học');pair('Từ ngày','Đến ngày')
p('Phạm vi lịch:  □ Hiện hành     □ Lịch sử',size=11)
table(['STT','Ngày','Giờ bắt đầu\nGiờ kết thúc','Lớp','Loại buổi','Phòng','Tình trạng'],[[str(i),'','','','','',''] for i in range(1,11)],[1.2,2.2,3.2,2.4,2.7,1.5,4.1],1.05,10.5)
h('GHI CHÚ VỀ LỊCH')
line('',110);line('',110)
small('Loại buổi: Buổi học / Thi cuối khóa. Tình trạng: Dự kiến / Đã tổ chức / Không tổ chức. Buổi không tổ chức có liên kết buổi thay thế, không tính thành hai buổi học. Ngày và giờ theo Việt Nam.')
p('Ngày lập bản lịch: ........ / ........ / ................',size=11,align=WD_ALIGN_PARAGRAPH.RIGHT)

start('BẢNG ĐIỂM CÁ NHÂN','BM-ĐT-03','Kết quả học tập đã được chốt và công bố')
line('Họ và tên học viên',78);pair('Mã học viên','Mã lớp');line('Khóa học',89);line('Chương trình đào tạo',74)
pair('Ngày công bố','Trạng thái')
small('Trạng thái bảng điểm: Đã chốt và công bố.')
table(['STT','Thành phần đánh giá','Trọng số\nphần trăm','Điểm\ntrên 100','Buổi liên quan nếu có'],[[str(i),'','','',''] for i in range(1,7)],[1.2,6.1,2.7,2.5,4.8],.95,11)
table(['ĐIỂM TỔNG KẾT TRÊN 100','KẾT QUẢ'],[['................................','□ Đạt      □ Chưa đạt']],[8.65,8.65],1.1,12)
h('NHẬN XÉT CỦA GIÁO VIÊN')
for _ in range(3):line('',110)
small('Thành phần và trọng số theo chương trình đào tạo. Tổng điểm làm tròn một lần đến hai chữ số thập phân; từ 50,00 trở lên là đạt. Kết quả chỉ được lập từ bảng điểm đã chốt.')
p('Ngày lập bản kết quả: ........ / ........ / ................',size=11,align=WD_ALIGN_PARAGRAPH.RIGHT)

start('BẢNG ĐIỂM DANH BUỔI HỌC','BM-ĐT-04','Điểm danh một buổi học của lớp')
pair('Mã lớp','Khóa học');line('Giáo viên phụ trách',75)
pair('Ngày học','Phòng');pair('Giờ bắt đầu','Giờ kết thúc')
table(['STT','Mã học viên','Họ và tên','Có mặt','Vắng'],[[str(i),'','','□','□'] for i in range(1,16)],[1.2,2.4,8.3,2.7,2.7],.70,11)
small('Đánh dấu một ô cho mỗi học viên đã được điểm danh. Cả hai ô để trống là chưa ghi nhận; không được hiểu là vắng. Không dùng bảng này cho buổi không tổ chức.')
sign('GIÁO VIÊN PHỤ TRÁCH')

start('PHIẾU XÁC NHẬN THANH TOÁN','BM-TC-01','Đối chiếu thông tin thanh toán học phí đã được Kế toán xác nhận')
pair('Số hóa đơn','Mã hồ sơ');line('Họ và tên học viên',78);line('Khóa học',89)
h('I THÔNG TIN THANH TOÁN')
line('Số tiền hóa đơn VND',74);line('Số tiền đã xác nhận VND',69)
p('Phương thức thanh toán:  □ Tiền mặt     □ Chuyển khoản')
line('Mã giao dịch nếu có',76);line('Thời điểm xác nhận',77);line('Người xác nhận',81)
p('Trạng thái hóa đơn sau xác nhận:  ĐÃ THANH TOÁN',True)
h('II ĐIỀU CHỈNH THÔNG TIN NẾU CÓ')
small('Chỉ ghi nội dung điều chỉnh phương thức hoặc mã giao dịch đã được xử lý trong hệ thống.')
line('Nội dung trước điều chỉnh',68);line('Nội dung sau điều chỉnh',70)
line('Lý do điều chỉnh',81);line('',110)
pair('Người điều chỉnh','Thời điểm điều chỉnh')
small('Thanh toán đủ một lần bằng số tiền hóa đơn. Phiếu ghi nhận thông tin đối chiếu; số tiền, hóa đơn và thời điểm xác nhận không thay đổi khi sửa phương thức hoặc mã giao dịch.')
sign('KẾ TOÁN XÁC NHẬN')

start('BÁO CÁO TỔNG HỢP HỌC PHÍ','BM-TC-02','Doanh thu và phải thu trong kỳ cùng công nợ hiện tại')
pair('Từ ngày','Đến ngày');line('Khóa học hoặc tất cả khóa',69)
line('Thời điểm lấy số liệu',73);p('Công nợ:  □ Tất cả công nợ     □ Chỉ quá hạn',size=11)
table(['Mã khóa','Tên khóa','Doanh thu\ntrong kỳ','Phải thu\ntrong kỳ','Công nợ\nhiện tại','Công nợ\nquá hạn'],[['','','','','',''] for _ in range(7)]+[['','TỔNG CỘNG','','','','']],[1.5,4.2,2.9,2.9,2.9,2.9],.9,10.5)
p('Đơn vị tính: VND nguyên đồng',size=11,align=WD_ALIGN_PARAGRAPH.RIGHT)
h('GHI CHÚ SỐ LIỆU')
small('Doanh thu theo thời điểm xác nhận thanh toán. Phải thu theo ngày phát hành hóa đơn, loại hóa đơn đã hủy. Công nợ là hóa đơn chưa thanh toán hiện tại; công nợ quá hạn là phần đã qua hạn thu.')
small('Tổng cộng tính trên toàn bộ tập dữ liệu đã lọc. Kỳ ngày không giới hạn công nợ hiện tại; không lấy phải thu trong kỳ trừ doanh thu trong kỳ để tính công nợ.')
sign('NGƯỜI LẬP BẢN BÁO CÁO')

start('HÓA ĐƠN HỌC PHÍ','BM-TC-03','Thông tin học phí của một hồ sơ đăng ký')
line('Số hóa đơn',88);pair('Ngày phát hành','Hạn thanh toán')
h('THÔNG TIN HỌC VIÊN VÀ KHÓA HỌC')
line('Họ và tên học viên',78);pair('Mã khóa học','Tên khóa học')
h('NỘI DUNG HỌC PHÍ')
table(['Nội dung','Số tiền VND'],[['Học phí khóa học đã đăng ký','................................']],[11.3,6],1.5,12)
table(['TỔNG HỌC PHÍ VND'],[['..........................................................................................']],[17.3],1.1,13)
h('TÌNH TRẠNG THANH TOÁN')
p('□ Chưa thanh toán     □ Đã thanh toán     □ Đã hủy',True)
line('Số ngày quá hạn nếu có',71)
h('LƯU Ý')
p('Hóa đơn chưa thanh toán không phải là xác nhận đã thu tiền.')
p('Số tiền được giữ theo học phí tại thời điểm phát hành. Hạn thu là ngày phát hành tại Việt Nam cộng 7 ngày lịch.')
p('Học viên thanh toán bên ngoài hệ thống; trạng thái thanh toán được cập nhật sau khi Kế toán xác nhận.')
small('Bản in đối chiếu từ thông tin hóa đơn mới nhất. Số hóa đơn do hệ thống cấp theo dạng HD-YYYY-NNNNNNNN. Quá hạn chỉ được tính với hóa đơn chưa thanh toán.')
p('Ngày lập bản in: ........ / ........ / ................',size=11,align=WD_ALIGN_PARAGRAPH.RIGHT)

start('HƯỚNG DẪN SỬ DỤNG BỘ MẪU','HD-BM','Trang hướng dẫn riêng không cần in kèm từng giấy tờ')
p('Bộ mẫu trình bày tám giấy tờ nghiệp vụ trên khổ A4. Các ô trống dành cho điền hoặc đưa dữ liệu hợp lệ vào trước khi in; mỗi mẫu nằm trên một trang. Dùng giấy A4, tỷ lệ in 100 phần trăm và in một mặt khi cần điền tay.')
table(['Trang','Biểu mẫu','Căn cứ sử dụng'],[
 ['1','Phiếu đăng ký khóa học','Một hồ sơ của Học viên; thông tin theo dõi lấy từ hệ thống.'],
 ['2','Danh sách lớp','Giáo viên hoặc Giáo vụ trong phạm vi lớp được phân công hoặc quản lý.'],
 ['3','Thời khóa biểu cá nhân','Lịch lớp của chính Học viên, phân biệt hiện hành và lịch sử.'],
 ['4','Bảng điểm cá nhân','Chỉ kết quả đã chốt của chính Học viên.'],
 ['5','Bảng điểm danh buổi học','Một buổi học hiện hành của lớp do Giáo viên phụ trách.'],
 ['6','Phiếu xác nhận thanh toán','Thông tin đã được Kế toán xác nhận; không thêm số phiếu thu riêng.'],
 ['7','Báo cáo tổng hợp học phí','Theo bộ lọc và số liệu tổng hợp; Giám đốc không phê duyệt.'],
 ['8','Hóa đơn học phí','Một hóa đơn của Học viên sở hữu, đọc mới trước khi in.'],
],[1.2,5.5,10.6],.8,11)
h('Quy tắc điền và ký')
p('Mã học viên, mã hồ sơ, số hóa đơn, trạng thái, thời điểm và kết quả phải đúng dữ liệu của hệ thống. Không tự chọn trạng thái hoặc tính lại tổng tài chính bằng các số đã làm tròn. Thông tin cá nhân sử dụng hồ sơ hiện tại; tiền hóa đơn giữ giá trị tại thời điểm phát hành.')
p('Các vùng ký dùng cho việc xác nhận hoặc đối chiếu bản giấy, không thay thế thao tác điện tử và không tạo thêm bước duyệt. Không thêm chữ ký Giám đốc phê duyệt, chữ ký số, dấu giả hoặc chứng từ thuế. Mẫu 06 là bản xác nhận đối chiếu, không chứng từ biên lai mới do ứng dụng phát hành.')
p('Điểm danh giấy phải được cập nhật bằng thao tác hợp lệ của Giáo viên nếu cần lưu vào hệ thống. Không dùng mẫu giấy để sửa điểm đã chốt, nhập điểm hoặc điểm danh ngoài cửa sổ cho phép, chuyển lớp, hoàn tiền hoặc thay đổi dữ liệu tài chính bất biến.')
small('Việc chuẩn bị mẫu giấy không bổ sung nút in hoặc quyền xuất vào phần mềm. Các quyền hiện có vẫn giữ: Học viên in hóa đơn của mình; Kế toán xuất CSV tài chính; các mẫu khác là tài liệu trình bày và đối chiếu riêng.')

d.save(out)
print(out)
