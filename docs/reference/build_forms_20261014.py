from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'docs/BIEU_MAU_NGHIEP_VU_THIET_KE_LAI_20261014.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(1.8)
sec.left_margin = sec.right_margin = Cm(1.8)
sec.header_distance = sec.footer_distance = Cm(.8)
sec.different_first_page_header_footer = False
doc.settings.odd_and_even_pages_header_footer = False
for border in list(doc.styles.element.xpath('.//w:pBdr')):
    border.getparent().remove(border)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3']:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0,0,0)
    style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Arial')
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(24)
doc.styles['Title'].paragraph_format.space_after = Pt(12)
doc.styles['Heading 1'].font.size = Pt(18)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(9)
doc.styles['Heading 2'].font.size = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(10)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(5)
header = sec.header.paragraphs[0]
header.text = 'TRUNG TÂM TIẾNG ANH  |  BIỂU MẪU NGHIỆP VỤ'
header.runs[0].font.size = Pt(8)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Trang ').font.size = Pt(8)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
doc.core_properties.title = 'Biểu mẫu nghiệp vụ thiết kế lại cho hệ thống quản lý trung tâm Anh ngữ'
doc.core_properties.subject = 'Biểu mẫu phù hợp với nghiệp vụ và quyền hiện hành'

def para(text, bold=False, size=None):
    p=doc.add_paragraph(); r=p.add_run(text); r.bold=bold
    if size:r.font.size=Pt(size)
    return p

def heading(text): doc.add_heading(text,2)

def fields(*pairs):
    for a,b in pairs:
        p=doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(8.7))
        p.add_run(a+'  ').bold=True; p.add_run('……………………')
        if b:p.add_run('\t'+b+'  ').bold=True;p.add_run('……………………')

def note(text):
    p=para(text,size=9)
    p.paragraph_format.space_before=Pt(3)
    return p

def table(headers, rows, widths, size=9.5):
    t=doc.add_table(rows=1, cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,width in zip(t.columns,widths):c.width=Cm(width)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for c,txt in zip(t.rows[0].cells,headers):c.text=txt
    for row in rows:
        cells=t.add_row().cells
        for c,txt in zip(cells,row):c.text=str(txt)
    for i,row in enumerate(t.rows):
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
        for j,c in enumerate(row.cells):
            c.width=Cm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            cb=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');cb.append(e)
            pr.append(cb)
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');m.append(e)
            pr.append(m)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EEF2' if i==0 else ('F6F8FA' if i%2==0 else 'FFFFFF'));pr.append(shade)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.05
                for r in p.runs:r.font.size=Pt(size);r.bold=i==0
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def start(n,title,code,role,kind):
    doc.add_page_break()
    doc.add_heading(f'{n} {title}',1)
    para(f'{code}  •  {role}',True,10)
    note(kind)

doc.add_paragraph('Biểu mẫu nghiệp vụ\nthiết kế lại',style='Title')
para('Hệ thống quản lý trung tâm Anh ngữ',True,13)
para('Bộ mẫu dùng để cập nhật phần Biểu mẫu của báo cáo, theo luồng và quyền hiện hành. Tám mẫu thống nhất cách ghi thông tin, phân biệt dữ liệu nhập với dữ liệu do hệ thống cung cấp và phản ánh đúng các bước đăng ký, đào tạo, học phí.')
table(['Mẫu','Tên biểu mẫu','Người sử dụng'],[
 ['01','Hồ sơ đăng ký khóa học','Học viên'],['02','Danh sách lớp','Giáo viên và Giáo vụ'],
 ['03','Thời khóa biểu cá nhân','Học viên'],['04','Bảng điểm cá nhân đã công bố','Học viên'],
 ['05','Điểm danh một buổi học','Giáo viên'],['06','Xác nhận thanh toán','Kế toán'],
 ['07','Báo cáo học phí','Kế toán và Giám đốc'],['08','Hóa đơn học phí','Học viên sở hữu'],
], [1.3,10.8,5.3],10)
heading('Quy ước sử dụng')
para('Dấu chấm và các dòng trống là vị trí điền hoặc trình bày dữ liệu, không phải dữ liệu mẫu đã phát sinh. Mã BM chỉ dùng để nhận diện biểu mẫu trong tài liệu. Mã hồ sơ, mã học viên nội bộ, số hóa đơn, thời điểm và kết quả do hệ thống cung cấp.')
para('Các biểu mẫu tra cứu và nhập liệu không tự tạo quyền in, xuất file hoặc phát hành chứng từ. Riêng hóa đơn của học viên được in qua trình duyệt; Kế toán được xuất CSV báo cáo tài chính. Mẫu 06 thay phiếu thu riêng bằng giao diện xác nhận thanh toán.')
para('Ngày theo dd/MM/yyyy, giờ theo HH:mm và giờ Việt Nam. Tiền nguyên VND; điểm theo thang 100. Thông tin cá nhân lấy từ hồ sơ hiện tại; học phí hóa đơn giữ giá trị tại thời điểm phát hành.')
heading('Căn cứ nghiệp vụ')
note('Yêu cầu nghiệp vụ chính thức, đặc tả chức năng và workflow đã chấp nhận; đối chiếu giao diện và dữ liệu hiện tại. Phiên bản biểu mẫu: 14/10/2026. File báo cáo gốc được giữ nguyên.')

start(1,'Hồ sơ đăng ký khóa học','BM-TS-01','Học viên','Biểu mẫu đăng ký và theo dõi hồ sơ của chính học viên. Một hồ sơ chọn một khóa học; không chọn lớp hoặc thu tiền tại bước tạo hồ sơ.')
heading('Thông tin học viên')
note('Lấy từ tài khoản và hồ sơ cá nhân. Học viên sửa thông tin qua chức năng hồ sơ riêng; email đăng nhập không đổi.')
fields(('Họ và tên','Mã học viên nội bộ'),('Ngày sinh','Số điện thoại'),('Email đăng nhập',None),('Địa chỉ không bắt buộc',None))
heading('Khóa học đăng ký')
fields(('Mã khóa học','Tên khóa học'),('Học phí hiện tại VND',None))
note('Học viên chọn một khóa đang hoạt động. Giá hiện tại là thông tin tham khảo; hóa đơn được phát hành khi Tuyển sinh duyệt hồ sơ, lấy học phí tại thời điểm đó.')
heading('Thông tin hồ sơ sau khi tạo')
fields(('Mã hồ sơ','Ngày đăng ký'),('Trạng thái hồ sơ','Trình độ đã xác định'))
table(['Lượt thi','Kết quả đã chấm trên 100','Tình trạng đặt ca'],[['1','',''],['2 nếu có','','']],[2.4,6.9,8.1])
fields(('Điểm tốt nhất trên 100',None),('Hóa đơn nếu đã phát hành','Lớp nếu đã được xếp'))
para('Lý do từ chối hoặc hủy khi có: ……………………………………………………………')
heading('Thao tác và điều kiện')
para('Tạo hồ sơ → chọn ca thi → xem kết quả và xác nhận nhập học → chờ Tuyển sinh duyệt → thanh toán ngoài hệ thống → chờ Giáo vụ xếp lớp.',size=10)
note('Tối đa hai lượt thi mỗi hồ sơ; vắng hoặc đổi ca không tạo lượt thứ ba. Không xác nhận khi còn lượt hai chờ kết quả. Học viên chỉ hủy trước xác nhận nhập học; đổi khóa cần hủy hợp lệ và tạo hồ sơ mới. Không sao chép điểm hoặc Band cũ.')

start(2,'Danh sách lớp','BM-ĐT-01','Giáo viên và Giáo vụ','Biểu mẫu tra cứu. Giáo viên chỉ xem lớp hiện được phân công; Học viên không được xem toàn bộ danh sách thành viên qua quyền này.')
heading('Thông tin lớp')
fields(('Mã lớp','Trạng thái lớp'),('Khóa học','Mức trình độ'),('Chương trình đào tạo',None),('Giáo viên phụ trách',None),('Số học viên đã xếp','Sức chứa tối đa'))
note('Sức chứa lớp trong 1–15. Phòng và giờ học thuộc từng buổi trong thời khóa biểu, không mặc định một phòng hoặc ca cố định cho cả lớp.')
heading('Danh sách thành viên')
table(['STT','Mã học viên','Họ và tên','Email','Số điện thoại'],[[str(i),'','','',''] for i in range(1,6)],[1.0,2.0,4.6,6.0,3.8])
note('Mã học viên là ID nội bộ. Tên, email và số điện thoại lấy từ thông tin hiện tại. Số dòng trống minh họa không giới hạn sĩ số hoặc thay cơ chế phân trang của hệ thống.')
heading('Thông tin bổ sung dành riêng cho Giáo vụ')
table(['Mã học viên','Mã hồ sơ','Band','Điểm đầu vào tốt nhất','Trạng thái hóa đơn'],[['','','','',''] for _ in range(2)],[2.2,2.2,2.0,5.2,5.8])
note('Khối này không thuộc danh sách Teacher. Giáo vụ chỉ xem trạng thái hóa đơn, không số tiền hoặc thông tin Payment.')
heading('Quy định')
para('Xếp lớp chỉ cho hồ sơ đã được duyệt và có hóa đơn PAID, đúng khóa/Band, lớp OPEN còn chỗ và chưa tới buổi học đầu, không trùng lịch. Tạo bảng điểm DRAFT cùng thao tác xếp lớp.')
para('Không có trường giới tính, thông tin phụ huynh, bảo lưu hoặc chuyển lớp. Không gỡ hoặc chuyển học viên sau khi đã xếp. Lớp hoàn thành và lịch sử vẫn được giữ theo quyền tra cứu.')
note('Danh sách này không phải phiếu phê duyệt hoặc biểu mẫu xuất file của nhân sự; không thêm ô chữ ký, đóng dấu.')

start(3,'Thời khóa biểu cá nhân','BM-ĐT-02','Học viên','Biểu mẫu tra cứu lớp và lịch của chính học viên. Dùng danh sách ngày cụ thể để thể hiện đúng giờ, phòng và buổi thay thế.')
heading('Phạm vi tra cứu')
fields(('Lớp hoặc khóa học',None),('Từ ngày','Đến ngày'))
para('Phạm vi:  □ Hiện hành    □ Lịch sử')
heading('Lịch học và thi cuối khóa')
table(['Ngày','Giờ bắt đầu và kết thúc','Lớp','Loại buổi','Phòng','Tình trạng'],[['','','','','',''] for _ in range(6)],[2.0,3.5,2.6,3.1,1.7,4.5])
note('Loại buổi: Buổi học LESSON hoặc Thi cuối khóa FINAL_EXAM. Tình trạng buổi: Dự kiến PLANNED, Đã tổ chức HELD hoặc Không tổ chức NOT_HELD. Hiện hành/Lịch sử được trình bày riêng với tình trạng buổi.')
heading('Chi tiết buổi được chọn')
fields(('Ngày học','Phòng'),('Giờ bắt đầu','Giờ kết thúc'),('Lớp','Giáo viên'),('Lịch hiện hành liên quan',None))
para('Nếu buổi không tổ chức: hiển thị lịch cũ và liên kết buổi thay thế cùng nghĩa vụ. Lịch cũ không xuất hiện như một buổi sắp tới hoặc bị đếm thêm vào tiến độ học.')
heading('Quy định sử dụng')
para('Học viên chỉ đọc lịch; không tự sửa, xếp hoặc chuyển lớp. Giáo vụ quản lý lịch theo cửa sổ thời gian và điều kiện nghiệp vụ. Không tự xác nhận đã tổ chức chỉ vì thời gian đã qua.')
para('Không giới hạn lịch vào ba khung sáng, chiều, tối cố định. Lịch tuần là cách hiển thị các buổi có ngày cụ thể; không có tự bỏ ngày lễ, tự sinh lịch bù hoặc gửi thông báo Email/Zalo.')
note('Ngày và giờ hiển thị theo Việt Nam. Thi đầu vào được theo dõi ở lịch đặt ca thi, không coi là LESSON hoặc FINAL_EXAM của lớp.')

start(4,'Bảng điểm cá nhân đã công bố','BM-ĐT-03','Học viên','Biểu mẫu chỉ xem kết quả của chính học viên sau khi lớp COMPLETED và bảng điểm FINALIZED. Không dùng làm form nhập điểm nháp của giáo viên.')
heading('Thông tin kết quả')
fields(('Mã lớp','Khóa học'),('Chương trình đào tạo',None),('Thời điểm công bố',None))
heading('Thành phần điểm')
table(['Thành phần','Loại','Trọng số phần trăm','Điểm trên 100','Buổi liên quan nếu có'],[['','','','',''] for _ in range(4)],[4.4,2.6,3.0,2.5,4.9])
note('Tên, số lượng, loại và trọng số lấy từ cấu hình chương trình đã khóa. GENERAL không gắn buổi; LESSON và FINAL_EXAM có lịch liên quan theo dữ liệu được phép xem. Không dùng hệ số giữa kỳ 1 và cuối kỳ 2 cố định.')
heading('Kết quả chính thức')
fields(('Điểm tổng kết trên 100',None))
para('Kết quả từ hệ thống:  □ Đạt    □ Chưa đạt')
note('Đạt khi điểm tổng kết FINALIZED từ 50,00 trở lên. Đây là kết quả chỉ đọc, không phải lựa chọn của học viên.')
heading('Nhận xét của giáo viên')
para('…………………………………………………………………………………………………')
para('…………………………………………………………………………………………………')
heading('Công thức và điều kiện công bố')
para('Điểm tổng kết = tổng của điểm thành phần × trọng số / 100. Tính bằng số thập phân, chỉ làm tròn tổng một lần theo half-up đến hai chữ số thập phân.')
para('Ví dụ về tổng trước làm tròn: 49,995 → 50,00; 49,994 → 49,99. Các số này minh họa phép tính tổng, không phải điểm đầu vào được phép nhập quá hai số lẻ.')
note('Chưa FINALIZED không được hiển thị cho Học viên; điểm thiếu không là 0. Không điều kiện chuyên cần bổ sung, xếp loại Giỏi/Khá hoặc mở lại điểm đã chốt. Không phát hành phiếu điểm ký đóng dấu hoặc chức năng tải PDF riêng.')

start(5,'Điểm danh một buổi học','BM-ĐT-04','Giáo viên','Biểu mẫu nhập liệu cho một LESSON hiện hành của lớp đang phụ trách. Lưu nhóm thay đổi của một buổi cùng thành công hoặc cùng hoàn tác.')
heading('Buổi học được chọn')
fields(('Mã lớp','Trạng thái lớp'),('Ngày học','Phòng'),('Giờ bắt đầu','Giờ kết thúc'),('Giáo viên phụ trách',None))
heading('Danh sách điểm danh')
table(['STT','Mã học viên','Họ và tên','Giá trị hiển thị hoặc lựa chọn'],[[str(i),'','','□ Chưa ghi   □ Có mặt   □ Vắng'] for i in range(1,6)],[1.0,2.2,5.4,8.8])
note('Mỗi học viên chỉ một giá trị. Có mặt lưu 1; Vắng lưu 0. Chưa ghi là chưa có dữ liệu, không phải một bản ghi vắng. Không mặc định cả lớp có mặt hoặc biến ô trống thành vắng.')
heading('Trạng thái thao tác')
para('Hiển thị: Chưa lưu / Đã lưu theo kết quả thực tế của lần thao tác.')
para('Thao tác:  [ Lưu điểm danh buổi ]    [ Bỏ thay đổi ]',True)
heading('Điều kiện ghi và sửa')
para('Được ghi hoặc sửa từ lúc LESSON bắt đầu đến trước khi lớp COMPLETED. Buổi phải hiện hành, không NOT_HELD; giáo viên vẫn đang được phân công. Sau COMPLETED chỉ xem, dù bảng điểm chưa gửi hoặc chưa chốt.')
para('Được lưu một phần danh sách có chủ đích; toàn bộ item gửi trong một lần phải hợp lệ. Thiếu điểm danh buổi học còn bắt buộc sẽ chặn hoàn thành lớp. Điểm danh không tự xác nhận buổi HELD.')
note('Không ký hiệu vắng có phép, vắng không phép hoặc đi muộn; không cảnh báo kỷ luật tự động và không ngưỡng vắng làm rớt môn. Số dòng minh họa không quyết định sĩ số lớp.')

start(6,'Xác nhận thanh toán','BM-TC-01','Kế toán','Biểu mẫu giao diện nghiệp vụ thay cho phiếu thu riêng. Không phát hành số phiếu PT, biên lai điện tử, chữ ký hoặc chứng từ để xuất và in.')
heading('Hóa đơn được đối chiếu')
note('Thông tin hệ thống cung cấp; không sửa số tiền hoặc liên kết hóa đơn.')
fields(('Số hóa đơn','Mã hồ sơ'),('Học viên','Khóa học'),('Số tiền phải thu VND',None),('Trạng thái hóa đơn','Hạn thanh toán'))
heading('Thông tin xác nhận')
para('Phương thức:  □ Tiền mặt CASH    □ Chuyển khoản TRANSFER')
fields(('Mã giao dịch',None))
note('Mã giao dịch bắt buộc khi chuyển khoản, tối đa 100 ký tự; tiền mặt có thể không có. Số tiền xác nhận bằng toàn bộ số tiền hóa đơn, không thu từng phần.')
para('Thao tác:  [ Chuẩn bị xác nhận ] → rà đối tượng và số tiền → [ Xác nhận thanh toán ]',True,10)
heading('Kết quả sau khi xác nhận thành công')
fields(('Trạng thái hóa đơn','Số tiền đã xác nhận VND'),('Người xác nhận',None),('Thời điểm xác nhận',None))
note('Thời điểm do server ghi; không nhập ngày thu thực tế riêng hoặc lùi ngày. Một Payment cho một Invoice, cùng giao dịch với chuyển Invoice PAID và nhật ký bắt buộc.')
heading('Chỉnh thông tin thanh toán đã xác nhận')
para('Chỉ sửa phương thức và mã giao dịch. Lý do sửa bắt buộc:')
para('…………………………………………………………………………………………………')
note('Giữ nguyên tiền, Invoice và thời điểm xác nhận; lưu cũ/mới/người/thời gian trong nhật ký. Không hủy, xóa, hoàn tiền hoặc tạo Payment thứ hai. Reset form hoặc in hóa đơn không xác nhận tiền đã nhận.')

start(7,'Báo cáo học phí','BM-TC-02','Kế toán và Giám đốc','Biểu mẫu kết quả tra cứu tổng hợp theo khóa. Kế toán được xuất CSV; Giám đốc chỉ xem tổng hợp, không xuất hoặc phê duyệt.')
heading('Bộ lọc đang áp dụng')
fields(('Từ ngày','Đến ngày'),('Khóa học hoặc tất cả',None),('Thời điểm truy vấn',None))
para('Công nợ hiện tại:  □ Tất cả công nợ    □ Chỉ quá hạn')
note('Kỳ ngày áp dụng cho doanh thu và phải thu trong kỳ. Công nợ và phần quá hạn là hiện tại, không giới hạn bằng kỳ ngày này.')
heading('Kết quả theo khóa học')
table(['Mã khóa','Tên khóa','Doanh thu trong kỳ VND','Phải thu trong kỳ VND','Công nợ hiện tại VND','Công nợ quá hạn VND'],[['','','','','',''] for _ in range(3)]+[['','Tổng toàn bộ tập lọc','','','','']],[1.5,3.7,3.05,3.05,3.05,3.05],9)
heading('Cơ sở tính các chỉ số')
para('Doanh thu: tổng tiền Payment được xác nhận trong kỳ. Phải thu: tổng Invoice UNPAID hoặc PAID phát hành trong kỳ, loại CANCELLED.')
para('Công nợ hiện tại: tổng Invoice UNPAID. Công nợ quá hạn: phần UNPAID có due date trước ngày Việt Nam hiện tại. Không lấy phải thu trong kỳ trừ doanh thu trong kỳ để tính công nợ.')
heading('Tra cứu và xuất kết quả')
para('Thao tác:  [ Áp dụng bộ lọc ]    [ Xuất CSV nếu là Kế toán ]',True)
note('Tổng tính trên toàn tập lọc, không cộng riêng trang đang xem. CSV truy vấn mới theo cùng bộ lọc và sáu cột; có thời điểm riêng nên có thể khác kết quả đã đọc trước đó.')
note('CSV UTF-8 BOM, dấu phẩy, CRLF, bảo vệ ô công thức; tối đa 10.000 dòng dữ liệu gồm dòng tổng và 5.000.000 byte gồm toàn bộ nội dung file. Vượt một giới hạn từ chối cả file, yêu cầu thu hẹp lọc.')
note('Không báo cáo chi phí, lương, bán tài liệu, marketing, lợi nhuận hoặc chữ ký Giám đốc phê duyệt. Giám đốc có báo cáo tuyển sinh và đào tạo riêng; không gộp chúng thành báo cáo thu chi.')

start(8,'Hóa đơn học phí','BM-TC-03','Học viên sở hữu hóa đơn','Mẫu đầu ra A4 qua chức năng in trình duyệt. Dữ liệu phải được đọc mới theo quyền trước khi mở bản in; không phải chứng từ thuế hoặc biên lai riêng.')
heading('Thông tin hóa đơn')
fields(('Số hóa đơn',None),('Học viên',None),('Mã khóa học','Tên khóa học'),('Thời điểm phát hành',None),('Hạn thanh toán',None))
heading('Học phí')
table(['Nội dung','Số tiền nguyên VND'],[['Học phí khóa học đã đăng ký','………………………………']],[11.0,6.4],11)
note('Một hóa đơn gắn một hồ sơ đăng ký. Số tiền được giữ theo học phí lúc phát hành; đổi giá khóa sau đó không đổi tiền hóa đơn cũ.')
heading('Trạng thái thanh toán')
para('Trạng thái từ hệ thống: ……………………………………………………………………',True)
para('Nếu UNPAID: Chưa thanh toán — không phải xác nhận đã thu tiền.')
para('Nếu PAID: Đã thanh toán. Nếu CANCELLED: Đã hủy.')
fields(('Số ngày quá hạn nếu có',None))
note('Quá hạn chỉ khi UNPAID và due date trước ngày hiện tại tại Việt Nam. Hạn thu là ngày phát hành Việt Nam +7 ngày lịch. Bấm in không chuyển hóa đơn PAID.')
heading('Quy định của bản in')
para('Bản in thực tế chỉ hiển thị dữ liệu hóa đơn được phép xem; ẩn menu và nút thao tác. Tên học viên lấy thông tin hiện tại; số tiền giữ snapshot đã phát hành.')
note('Số hóa đơn thực tế có định dạng HD-YYYY-NNNNNNNN do server cấp. Không dòng phí tài liệu, quantity, discount, VAT, mã số thuế, tiền bằng chữ, chữ ký hoặc đóng dấu giả. Không đưa Payment method/reference hoặc người xác nhận vào bản in Student.')
note('Không có API xuất PDF riêng hoặc quyền in cho nhân sự. Người dùng có thể dùng lựa chọn lưu PDF của trình duyệt nếu trình duyệt hỗ trợ.')

doc.save(OUT)
print(OUT)
