from pathlib import Path
import json,math,html,zipfile,hashlib,re
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor
import pypdfium2 as pdfium
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
DIR=ROOT/'docs/reference/usecase-update-20261014';DIR.mkdir(exist_ok=True)
ASSETS=DIR/'diagrams';ASSETS.mkdir(exist_ok=True)
pdfmetrics.registerFont(TTFont('UCArial',r'C:\Windows\Fonts\arial.ttf'))
pdfmetrics.registerFont(TTFont('UCArialBold',r'C:\Windows\Fonts\arialbd.ttf'))
# Each association grants only the scope specified in the description, never a new role.
diagrams=[]
def D(code,title,actors,left,right,edges,includes,scope,guards):
 diagrams.append(dict(code=code,title=title,actors=actors,left=left,right=right,edges=edges,includes=includes,scope=scope,guards=guards))
A=lambda key,label,side,y:dict(key=key,label=label,side=side,y=y)
D('UC00','Tổng quan phạm vi hệ thống',
 [A('s','Học viên','L',650),A('ad','Tuyển sinh','L',420),A('t','Giáo viên','L',190),A('admin','Quản trị viên','R',650),A('aa','Giáo vụ','R',490),A('ac','Kế toán','R',320),A('di','Giám đốc','R',150)],
 [('auth','Tài khoản và phiên dùng chung'),('reg','Hồ sơ đăng ký và thi đầu vào'),('own','Lớp lịch và kết quả cá nhân'),('teach','Giảng dạy và điểm danh')],
 [('staff','Quản trị tài khoản và nhật ký'),('acad','Cấu hình lớp và lịch'),('finance','Hóa đơn và thanh toán'),('reports','Báo cáo tổng hợp')],
 [('s','auth'),('ad','auth'),('t','auth'),('admin','auth'),('aa','auth'),('ac','auth'),('di','auth'),('s','reg'),('ad','reg'),('t','reg'),('aa','reg'),('s','own'),('t','teach'),('admin','staff'),('aa','acad'),('ac','finance'),('s','finance'),('di','reports'),('ac','reports')],[],
 'Hình tổng quan gom các mục tiêu thành nhóm để giới thiệu phạm vi; các hình UC01–UC09 xác định quyền cụ thể. Bảy tác nhân bên ngoài tương ứng bảy vai trò cố định.',
 'Đường nối thể hiện tham gia theo phạm vi, không phải quyền quản lý toàn bộ nhóm. “Hệ thống” không được vẽ thành tác nhân bên ngoài; không có phê duyệt của Giám đốc.')
D('UC01','Tài khoản và phiên đăng nhập',
 [A('u','Người dùng\n7 vai trò','L',600),A('s','Học viên','L',245),A('admin','Quản trị viên','R',420)],
 [('login','Đăng nhập'),('logout','Đăng xuất phiên hiện tại'),('pass','Tự đổi mật khẩu'),('signup','Tự đăng ký tài khoản Học viên'),('profile','Sửa hồ sơ cá nhân của mình')],
 [('create','Cấp tài khoản nhân sự'),('reset','Đặt lại mật khẩu tạm'),('lock','Khóa hoặc mở khóa tài khoản'),('disable','Ngừng hoặc khôi phục tài khoản'),('revoke','Thu hồi các phiên liên quan')],
 [('u','login'),('u','logout'),('u','pass'),('s','signup'),('s','profile'),('admin','create'),('admin','reset'),('admin','lock'),('admin','disable')],
 [('pass','revoke'),('reset','revoke')],
 '“Người dùng 7 vai trò” là ký hiệu rút gọn cho quyền dùng chung, không phải vai trò thứ tám. Student tự tạo tài khoản; Admin cấp nhân sự và quản lý trạng thái tài khoản.',
 'Email bất biến; tự đổi mật khẩu cần mật khẩu hiện tại. Reset/đổi mật khẩu/khóa/ngừng dùng thu hồi mọi phiên; mở khóa/khôi phục không phục hồi phiên cũ. Quan hệ include ở hình thể hiện hai lệnh luôn thu hồi; nhánh khóa/ngừng cũng thu hồi theo điều kiện thao tác. Bảo vệ Admin tự khóa và Admin hoạt động cuối cùng.')
D('UC02','Hồ sơ đăng ký và quyết định tuyển sinh',
 [A('s','Học viên','L',435),A('ad','Tuyển sinh','R',470)],
 [('courses','Xem khóa học đang hoạt động'),('create','Tạo hồ sơ đăng ký khóa học'),('own','Xem hồ sơ của mình'),('cancelown','Hủy hồ sơ trước xác nhận'),('confirm','Xác nhận nhập học')],
 [('search','Tra cứu và kiểm tra hồ sơ'),('approve','Duyệt hồ sơ nhập học'),('reject','Từ chối hồ sơ có lý do'),('cancel','Hủy hồ sơ theo quyền Tuyển sinh'),('invoice','Phát hành một hóa đơn học phí')],
 [('s','courses'),('s','create'),('s','own'),('s','cancelown'),('s','confirm'),('ad','search'),('ad','approve'),('ad','reject'),('ad','cancel')],
 [('approve','invoice')],
 'Student tự đăng ký và xác nhận sau kết quả đầu vào; Admissions tra cứu và quyết định hồ sơ đã CONFIRMED. Phát hành Invoice là xử lý bắt buộc bên trong thao tác duyệt.',
 'Một hồ sơ chưa kết thúc cho mỗi Student/Course. Confirmation không tạo Invoice. Duyệt tạo ADMITTED + Invoice + audit trong một giao dịch; hủy hợp lệ xử lý Invoice UNPAID và booking tương lai, kể cả ca CLOSED, giữ lịch sử. Không ghi danh/đặt ca thay hoặc hủy PAID.')
D('UC03','Thi đầu vào và xác định trình độ',
 [A('s','Học viên','L',565),A('aa','Giáo vụ','L',205),A('t','Giáo viên chấm','R',560)],
 [('slots','Xem ca phù hợp và lịch của mình'),('booking','Đặt đổi hủy hoặc đặt lại ca thi'),('results','Xem kết quả từng lượt và Band'),('session','Tạo mở đóng hủy hoặc sửa ca thi'),('grader','Phân công hoặc đổi người chấm'),('follow','Theo dõi kết quả và tiến độ chấm')],
 [('assigned','Xem ca và danh sách được giao'),('grade','Nhập điểm đầu vào'),('absent','Ghi nhận vắng thi'),('correct','Sửa điểm đầu vào có lý do'),('band','Tính lại điểm tốt nhất và Band')],
 [('s','slots'),('s','booking'),('s','results'),('aa','session'),('aa','grader'),('aa','follow'),('t','assigned'),('t','grade'),('t','absent'),('t','correct')],
 [('grade','band'),('correct','band')],
 'Booking là lịch sử đặt ca; logical attempt là lượt thi, tối đa hai mỗi Registration. Thay đổi hoặc đặt lại booking không tạo lượt thứ ba. Band do hệ thống suy từ điểm Teacher nhập.',
 'Student chỉ thao tác booking OPEN trước start; lượt hai cần điểm lượt một và trước CONFIRMED. Initial score/ABSENT từ session end; chỉ original grader còn assigned sửa điểm, trước confirmation và có reason. Vắng không là 0. Hòa chọn lượt một; bộ ngưỡng khóa từ completed result đầu tiên.')
D('UC04','Hóa đơn và xác nhận thanh toán',
 [A('s','Học viên','L',535),A('ac','Kế toán','R',470)],
 [('own','Xem hóa đơn và hướng dẫn thanh toán'),('print','In hóa đơn của mình'),('fresh','Đọc mới hóa đơn đúng quyền')],
 [('search','Tra cứu hóa đơn và công nợ'),('pay','Xác nhận thanh toán đủ một lần'),('correct','Sửa phương thức hoặc mã giao dịch'),('paid','Lưu Payment và chuyển Invoice PAID')],
 [('s','own'),('s','print'),('ac','search'),('ac','pay'),('ac','correct')],
 [('print','fresh'),('pay','paid')],
 'Student xem/in Invoice riêng và trả tiền ngoài ứng dụng. Accountant xác nhận tiền đã đối chiếu và được sửa metadata Payment theo điều kiện; không tạo lại hóa đơn khi thu tiền.',
 'Invoice phát hành từ duyệt tuyển sinh ở UC02. Payment bằng toàn bộ amount, một lần, CASH/TRANSFER; chuyển khoản cần reference. Chỉ sửa method/reference có reason/audit, giữ money/link/time. Không gateway, trả góp, refund, biên lai riêng; overdue chỉ cảnh báo.')
D('UC05','Cấu hình khóa và chương trình đào tạo',
 [A('aa','Giáo vụ','L',435)],
 [('course','Tạo sửa ngừng hoặc khôi phục khóa'),('band','Cấu hình các ngưỡng Band'),('program','Tạo chương trình theo khóa và Band')],
 [('config','Cấu hình số buổi thành phần trọng số'),('activate','Kiểm tra và kích hoạt chương trình'),('state','Chuyển nháp ngừng hoặc khôi phục')],
 [('aa','course'),('aa','band'),('aa','program'),('aa','config'),('aa','activate'),('aa','state')],[],
 'AcademicAffairs quản lý Course, Band và TrainingProgram. Program thuộc cặp Course/Band duy nhất, có expected LESSON count và các ScoreComponentConfig.',
 'ACTIVE cần count dương, total weight 100 và đúng một FINAL_EXAM component. First Class use khóa Course/Band/components/kind/weights/count, kể cả Class CANCELLED. Giá mới chỉ cho Invoice sau đó; first entry result khóa bộ ngưỡng Band. Không version né khóa hay phê duyệt Director.')
D('UC06','Lớp xếp lớp và quản lý lịch',
 [A('aa','Giáo vụ','L',475),A('s','Học viên','R',170)],
 [('class','Tạo và quản lý sức chứa lớp'),('lifecycle','Mở bắt đầu hoàn thành hoặc hủy lớp'),('teacher','Phân công hoặc đổi giáo viên lớp'),('place','Xếp học viên đủ điều kiện vào lớp'),('weekly','Lập lịch tuần hoặc từng buổi')],
 [('move','Dời lịch hoặc thay buổi không tổ chức'),('held','Xác nhận buổi đã tổ chức'),('draft','Tạo bảng điểm nháp khi xếp lớp'),('own','Xem lớp và lịch cá nhân')],
 [('aa','class'),('aa','lifecycle'),('aa','teacher'),('aa','place'),('aa','weekly'),('aa','move'),('aa','held'),('s','own')],
 [('place','draft')],
 'Giáo vụ quản lý Class, assignments và từng Schedule cụ thể. Xếp lớp tạo một assignment và GradeReport DRAFT cùng giao dịch; Student chỉ xem lớp/lịch của mình.',
 'Chỉ OPEN trước buổi đầu, ADMITTED + PAID, đúng Course/Band, còn chỗ và không trùng lịch. Không gỡ/chuyển sau xếp. Lịch tuần all-or-none; lịch mới/đích dời ở tương lai, replacement giữ cùng nghĩa vụ; HELD do Giáo vụ sau end/ONGOING. COMPLETE cần mọi nghĩa vụ HELD và đủ LESSON Attendance, chưa yêu cầu chốt điểm.')
D('UC07','Giảng dạy và bảng điểm cả lớp',
 [A('t','Giáo viên','L',450),A('aa','Giáo vụ','R',620),A('s','Học viên','R',155)],
 [('classes','Xem lớp danh sách và lịch được giao'),('attendance','Ghi hoặc sửa điểm danh một buổi'),('draft','Lưu điểm nháp và nhận xét một học viên'),('submit','Gửi bảng điểm cả lớp'),('resubmit','Sửa và gửi lại lớp bị trả')],
 [('review','Xem bảng điểm theo quyền Giáo vụ'),('return','Trả lại cả lớp có lý do'),('finalize','Chốt và công bố cả lớp'),('total','Tính và lưu điểm tổng kết chính thức'),('published','Xem kết quả cá nhân đã chốt')],
 [('t','classes'),('t','attendance'),('t','draft'),('t','submit'),('t','resubmit'),('aa','review'),('aa','return'),('aa','finalize'),('s','published')],
 [('finalize','total')],
 'Teacher lưu Attendance theo một LESSON và draft theo một Student. Submit/resubmit/return/finalize là các lệnh cho toàn Class. Student chỉ xem own FINALIZED.',
 'Attendance từ LESSON start đến trước COMPLETED; draft ONGOING/COMPLETED và DRAFT/RETURNED. Submit khi COMPLETED/đủ mọi điểm/không sửa chưa lưu; return/finalize toàn bộ SUBMITTED, all-or-none. Thang 100, half-up tổng một lần hai chữ số; FINALIZED ≥50 đạt, bất biến; không ghi final_score trước chốt.')
D('UC08','Báo cáo tổng hợp và CSV tài chính',
 [A('di','Giám đốc','L',440),A('ac','Kế toán','R',525)],
 [('admissions','Xem báo cáo tuyển sinh tổng hợp'),('training','Xem báo cáo đào tạo và tỷ lệ đạt'),('finance','Xem báo cáo tài chính tổng hợp')],
 [('accountant','Xem báo cáo học phí theo quyền'),('csv','Xuất CSV tài chính toàn tập lọc'),('fresh','Truy vấn mới kết quả tài chính')],
 [('di','admissions'),('di','training'),('di','finance'),('ac','accountant'),('ac','csv')],
 [('csv','fresh')],
 'Director chỉ đọc ba nhóm aggregates không định danh, không export/phê duyệt. Accountant chỉ finance, có CSV toàn tập lọc; xuất chạy query mới, không chỉ trang đang xem.',
 'Doanh thu theo Payment confirmation; phải thu theo Invoice issue, loại CANCELLED; debt UNPAID hiện tại không kỳ ngày. Distinct Student toàn tập không cộng nhóm; tiến độ HELD obligations; pass rate FINALIZED-date, mẫu số rỗng là chưa có dữ liệu. CSV six columns/footer, caps và formula safeguards giữ nguyên.')
D('UC09','Tra cứu nhật ký hoạt động',
 [A('admin','Quản trị viên','L',440)],
 [('list','Tra cứu nhật ký có bộ lọc'),('detail','Xem chi tiết sự kiện được phép'),('group','Xem các sự kiện cùng thao tác')],[],
 [('admin','list'),('admin','detail'),('admin','group')],[],
 'Admin xem AuditLog newest-first/paginated, event detail và liên kết operation_id. Quyền đọc log không mở thêm quyền truy cập đối tượng nghiệp vụ.',
 'NV39 là xử lý nội bộ xuyên các workflow: mandatory mutation audit atomic/rollback; auth/logout/denial security-only best-effort. Không audit mọi read, lưu secrets hoặc raw payload. Mandatory history giữ; pure security/IP operational expiry 90 ngày. Không sửa/xóa/undo/export log.')

def lines(text,font,size,max_width):
 out=[]
 for segment in text.split('\n'):
  line=''
  for word in segment.split():
   candidate=(line+' '+word).strip()
   if line and pdfmetrics.stringWidth(candidate,font,size)>max_width:out.append(line);line=word
   else:line=candidate
  out.append(line)
 return out

def diagram_draw(item):
 W,H=1260,850;path=ASSETS/(item['code']+'.pdf');c=canvas.Canvas(str(path),pagesize=(W,H));c.setTitle(item['title'])
 svg=[]
 def ln(x1,y1,x2,y2,dash=False):
  c.setStrokeColor(HexColor('#48515a'));c.setLineWidth(1.2);c.setDash(5,4) if dash else c.setDash();c.line(x1,y1,x2,y2)
  svg.append(f'<line x1="{x1}" y1="{H-y1}" x2="{x2}" y2="{H-y2}" stroke="#48515a" stroke-width="1.2"'+(' stroke-dasharray="5 4"' if dash else '')+'/>')
 def text(x,y,value,size=16,bold=False,width=240):
  font='UCArialBold' if bold else 'UCArial';ls=lines(value,font,size,width);start=y+(len(ls)-1)*(size+3)/2
  c.setFillColor(HexColor('#15191d'));c.setFont(font,size)
  for i,s in enumerate(ls):
   yy=start-i*(size+3);c.drawCentredString(x,yy-size*.33,s)
   svg.append(f'<text x="{x}" y="{H-yy+size*.33}" font-family="Arial" font-size="{size}" font-weight="{"bold" if bold else "normal"}" text-anchor="middle" fill="#15191d">{html.escape(s)}</text>')
 c.setStrokeColor(HexColor('#6b7280'));c.setLineWidth(1.2);c.rect(230,75,800,690,fill=0)
 svg.append('<rect x="230" y="85" width="800" height="690" fill="white" stroke="#6b7280" stroke-width="1.2"/>')
 text(630,797,item['code']+'  '+item['title'],23,True,1100)
 text(630,739,'Hệ thống quản lý trung tâm Anh ngữ',15,True,700)
 positions={}
 for col,x in [('left',440 if item['right'] else 630),('right',820)]:
  nodes=item[col];n=len(nodes)
  for i,(key,label) in enumerate(nodes):
   y=680-i*(530/max(n-1,1)) if n>1 else 440
   positions[key]=(x,y,label)
 # Associations are drawn first, then ovals, preventing lines from crossing labels.
 for actor,node in item['edges']:
  ac=next(a for a in item['actors'] if a['key']==actor);x1=115 if ac['side']=='L' else 1145;y1=ac['y'];x,y,_=positions[node]
  ln(x1+(25 if ac['side']=='L' else -25),y1,x+(-135 if ac['side']=='L' else 135),y)
 for relation_index,(source,target) in enumerate(item['includes']):
  x1,y1,_=positions[source];x2,y2,_=positions[target]
  if x1==x2:
   side=-1 if x1>630 else 1
   a=(x1+side*135,y1);b=(x2+side*135,y2);gutter=620+relation_index*12
   ln(*a,gutter,y1,True);ln(gutter,y1,gutter,y2,True);ln(gutter,y2,*b,True)
   previous=(gutter,y2);mx,my=630,(y1+y2)/2
  else:
   dx=x2-x1;dy=y2-y1;norm=math.sqrt((dx/135)**2+(dy/36)**2)
   a=(x1+dx/norm,y1+dy/norm);b=(x2-dx/norm,y2-dy/norm)
   ln(*a,*b,True);previous=a;mx,my=(a[0]+b[0])/2,(a[1]+b[1])/2
  angle=math.atan2(b[1]-previous[1],b[0]-previous[0]);length=12
  for ang in [angle+.42,angle-.42]:ln(b[0],b[1],b[0]-length*math.cos(ang),b[1]-length*math.sin(ang))
  c.setFillColor(HexColor('#FFFFFF'));c.rect(mx-44,my-10,88,20,stroke=0,fill=1)
  svg.append(f'<rect x="{mx-44}" y="{H-my-10}" width="88" height="20" fill="white"/>')
  text(mx,my,'«include»',13,width=100)
 for key,(x,y,label) in positions.items():
  c.setDash();c.setStrokeColor(HexColor('#4b5563'));c.setFillColor(HexColor('#f4f7f9'));c.ellipse(x-135,y-36,x+135,y+36,stroke=1,fill=1)
  svg.append(f'<ellipse cx="{x}" cy="{H-y}" rx="135" ry="36" fill="#f4f7f9" stroke="#4b5563"/>');text(x,y,label,16,width=225)
 for actor in item['actors']:
  x=105 if actor['side']=='L' else 1155;y=actor['y']
  c.setStrokeColor(HexColor('#333333'));c.setFillColor(HexColor('#FFFFFF'));c.circle(x,y+42,13,stroke=1,fill=1)
  svg.append(f'<circle cx="{x}" cy="{H-y-42}" r="13" fill="white" stroke="#333333"/>')
  for xy in [(x,y+29,x,y-15),(x-24,y+12,x+24,y+12),(x,y-15,x-23,y-47),(x,y-15,x+23,y-47)]:ln(*xy)
  text(x,y-77,actor['label'],17,True,195)
 legend='Đăng nhập đổi mật khẩu và đăng xuất áp dụng cho cả 7 vai trò — chi tiết ở UC01' if item['code']=='UC00' else 'Đường liền: tham gia theo quyền   •   «include»: xử lý bắt buộc của use case nguồn'
 text(630,37,legend,13,width=1100)
 c.save()
 (ASSETS/(item['code']+'.svg')).write_text('<svg xmlns="http://www.w3.org/2000/svg" width="1260" height="850" viewBox="0 0 1260 850">'+''.join(svg)+'</svg>',encoding='utf-8')
 with pdfium.PdfDocument(str(path)) as pdf:
  page=pdf[0];bitmap=page.render(scale=1.6);bitmap.to_pil().save(ASSETS/(item['code']+'.png'));bitmap.close();page.close()
 pu=['@startuml','left to right direction','skinparam monochrome true','skinparam shadowing false']
 for actor in item['actors']:pu.append(f'actor "{actor["label"].replace(chr(10),chr(92)+"n")}" as A_{actor["key"]}')
 pu.append('rectangle "Hệ thống quản lý trung tâm Anh ngữ" {')
 for key,(_,_,label) in positions.items():pu.append(f'usecase "{label}" as U_{key}')
 pu.append('}')
 for actor,node in item['edges']:pu.append(f'A_{actor} -- U_{node}')
 for source,target in item['includes']:pu.append(f'U_{source} ..> U_{target} : <<include>>')
 pu.append('@enduml');(ASSETS/(item['code']+'.puml')).write_text('\n'.join(pu),encoding='utf-8')

# Keep the overview readable: one functional group per role, with shared auth in the caption.
diagrams[0]['left']=[('student','Tự đăng ký và theo dõi việc học'),('admissions','Xử lý hồ sơ tuyển sinh'),('teacher','Thi đầu vào và giảng dạy')]
diagrams[0]['right']=[('staff','Quản trị tài khoản và nhật ký'),('acad','Cấu hình lớp lịch và duyệt điểm'),('finance','Học phí thanh toán và CSV'),('reports','Báo cáo quản lý tổng hợp')]
diagrams[0]['edges']=[('s','student'),('ad','admissions'),('t','teacher'),('admin','staff'),('aa','acad'),('ac','finance'),('di','reports')]
diagrams[5]['left']=diagrams[5]['left']+diagrams[5]['right']
diagrams[5]['right']=[]
diagrams[2]['guards']+=' Invoice giữ học phí tại phát hành và hạn ngày Việt Nam +7 ngày lịch.'
diagrams[3]['guards']+=' Band chọn min_score lớn nhất không vượt điểm tốt nhất; seed ngưỡng 0/25/45/65/85.'
for item in diagrams:diagram_draw(item)
(DIR/'model.json').write_text(json.dumps(diagrams,ensure_ascii=False,indent=2),encoding='utf-8')

# Trace retains all 37 accepted source NV codes; UI/scope refinements do not introduce new NV.
trace=[
 ('01','Tự đăng ký và đăng nhập','UC01','F01–02','W01–02'),('02','Xem khóa tạo và hủy hồ sơ','UC02','F04–05','W04–05'),('03','Đặt đổi hủy và đặt lại ca thi','UC03','F06–07','W07'),('04','Xem kết quả đầu vào và Band','UC03','F08','W08'),('05','Xác nhận nhập học','UC02','F09','W09'),('06','Xem và in hóa đơn riêng','UC04','F10','W11'),('07','Xem lớp và lịch cá nhân','UC06','F11','W15–16'),('08','Xem bảng điểm đã chốt','UC07','F12','W19'),('09','Tra cứu kiểm tra hồ sơ','UC02','F13–14','W10'),('12','Duyệt từ chối hủy tuyển sinh','UC02','F14','W10–11'),('13','Khóa Program Band và cấu hình','UC05','F18','W13'),('14','Tổ chức và quản lý ca thi','UC03','F16–17','W06'),('15','Theo dõi entry result và Band','UC03','F16–17','W06–08'),('16','Tạo và quản lý vòng đời lớp','UC06','F19–20','W14'),('17','Phân công và đổi Teacher lớp','UC06','F20','W14'),('18','Xếp Student đủ điều kiện','UC06','F22','W15'),('19','Lịch tuần dời thay xác nhận HELD','UC06','F21','W16'),('20','Trả lại và chốt điểm cả lớp','UC07','F23','W19'),('21','Xem ca được phân công','UC03','F25','W06'),('22','Chấm sửa entry score ghi vắng','UC03','F25','W08'),('23','Xem lớp roster và lịch dạy','UC07','F24','W14–16'),('24','Điểm danh một LESSON','UC07','F26','W17'),('25','Lưu nháp và gửi bảng điểm','UC07','F27','W18'),('26','Sửa và gửi lại sau RETURNED','UC07','F27','W18'),('27','Tra cứu Invoice nghiệp vụ','UC04','F28','W11'),('28','Xác nhận và sửa Payment metadata','UC04','F29','W12'),('29','Finance report và CSV','UC08','F30','W20'),('30','Director admissions aggregates','UC08','F31','W20'),('31','Director training aggregates','UC08','F31','W20'),('32','Director finance aggregates','UC08','F31','W20'),('33','Cấp staff reset temporary password','UC01','F32','W01–03'),('34','Lock unlock disable restore','UC01','F32','W03'),('35','Audit viewer và operation grouping','UC09','F33','W21'),('36','Best completed score và Band','UC03','F08/F25','W08'),('37','Invoice trong admission approval','UC02','F14/F10','W10–11'),('38','Draft total và finalized total','UC07','F27/F23/F12','W18–19'),('39','Mandatory audit và security events','UC01–09 quy tắc chung','F33 khi đọc','W21')]
assert {int(x[0]) for x in trace}==set(range(1,40))-{10,11}
changes=[
 ('Tổng thể','Nhóm chức năng quá rộng; quyền dùng chung chưa đầy đủ.','UC00 tổng quan, UC01–09 chi tiết; mọi liên kết giữ scope của actor.'),
 ('Tài khoản','Chỉ Student đăng nhập; thiếu logout/password/profile và các thao tác Admin.','UC01 thêm chức năng chung, restricted credentials và thu hồi phiên.'),
 ('Tuyển sinh','Có hỗ trợ ghi danh; thiếu booking/test/confirmation/reject/cancel.','UC02 loại làm thay; UC03 tách entry test, Teacher chấm và Band.'),
 ('Học phí','Student “thanh toán hóa đơn”; payment include tạo Invoice.','UC04 Student xem/in/hướng dẫn; Invoice include từ approval ở UC02, không payment.'),
 ('Đào tạo','Thiếu Program/Band/entry sessions/lifecycle/replacement; có actor không nối.','UC05 cấu hình; UC06 Class/placement/schedules; UC03 ca thi; actor chỉ xuất hiện khi tham gia.'),
 ('Giảng dạy','Chỉ nhập điểm include tính điểm, thiếu whole-class review/publication.','UC07 tách draft total với official total; submit/return/finalize cả lớp và Student published-only.'),
 ('Báo cáo','Director có ba nhóm nhưng thiếu scope, Accountant và CSV.','UC08 giữ Director aggregate-only; CSV Accountant/query mới; UC09 tách Admin audit.')]

doc=Document();sec=doc.sections[0];sec.page_width=Cm(29.7);sec.page_height=Cm(21);sec.left_margin=sec.right_margin=Cm(1.5);sec.top_margin=Cm(1.3);sec.bottom_margin=Cm(1.2)
for b in list(doc.styles.element.xpath('.//w:pBdr')):b.getparent().remove(b)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0);st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Times New Roman')
 for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:st._element.get_or_add_rPr().rFonts.attrib.pop(qn('w:'+attr),None)
doc.styles['Normal'].font.size=Pt(11);doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Title'].font.size=Pt(22);doc.styles['Heading 1'].font.size=Pt(16);doc.styles['Heading 2'].font.size=Pt(12)
doc.core_properties.title='Use case cập nhật theo dự án quản lý trung tâm Anh ngữ'
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;footer.add_run('Trang ').font.size=Pt(8);field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
def para(t):doc.add_paragraph(t)
def table(head,rows,widths):
 t=doc.add_table(rows=1,cols=len(head));t.autofit=False
 for col,width in zip(t.columns,widths):col.width=Cm(width)
 for j,txt in enumerate(head):t.rows[0].cells[j].text=txt
 for vals in rows:
  cells=t.add_row().cells
  for c,txt in zip(cells,vals):c.text=str(txt)
 for i,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if i==0:pr.append(OxmlElement('w:tblHeader'))
  for j,c in enumerate(row.cells):
   c.width=Cm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   cp=c._tc.get_or_add_tcPr();bd=OxmlElement('w:tcBorders')
   for e in ['top','left','bottom','right']:
    b=OxmlElement('w:'+e);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');bd.append(b)
   cp.append(bd);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EAEFF3' if i==0 else 'FFFFFF');cp.append(sh)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(5)
    for r in p.runs:r.font.size=Pt(10.5);r.bold=i==0
 return t

doc.add_paragraph('Use case cập nhật theo dự án',style='Title')
para('Hệ thống quản lý trung tâm Anh ngữ   •   Phiên bản 14/10/2026')
para('Bộ tài liệu thay thế nội dung các sơ đồ use case cũ trong phần V của báo cáo. Các hình cập nhật đúng chức năng, quyền, điều kiện nghiệp vụ và xử lý nội bộ hiện hành; file báo cáo gốc được giữ nguyên.')
doc.add_heading('Phạm vi và quy ước',2)
para('Bảy vai trò cố định: Học viên, Tuyển sinh, Giáo vụ, Giáo viên, Kế toán, Giám đốc, Quản trị viên. Không vẽ phần mềm, database hoặc “Hệ thống” thành tác nhân thứ tám. UC00 là bản đồ nhóm chức năng; UC01–UC09 là các sơ đồ chi tiết. Mã UC chỉ định danh hình, không thay mã NV/F/W hoặc tạo chức năng mới.')
para('Đường liên kết actor–use case thể hiện tham gia trong phạm vi quyền, không mang thứ tự xử lý. Đường nét đứt «include» đi từ use case nguồn đến xử lý bắt buộc được gọi. Các thao tác lựa chọn như hủy, từ chối, sửa, trả lại được vẽ thành use case riêng; không dùng «extend» để thay một bước tiếp theo của workflow.')
para('Đăng nhập/phiên hợp lệ là tiền điều kiện của nghiệp vụ, không include vào mọi thao tác. Một use case có thể gộp các lựa chọn cùng mục tiêu, ví dụ đặt/đổi/hủy/đặt lại booking; mô tả điều kiện chỉ rõ từng nhánh. Xử lý nội bộ như Band, Invoice và tổng điểm được đặt trong ranh giới hệ thống.')
doc.add_heading('Các ràng buộc xuyên suốt',2)
para('Quyền và ownership được kiểm tra ở backend; Teacher chỉ ca/lớp hiện assigned, Student chỉ dữ liệu riêng, Director chỉ tổng hợp không định danh. Preview không giữ chỗ; server kiểm lại state/time/resources. Các lệnh quan trọng giữ replay safety và audit bắt buộc nguyên tử; mất phản hồi cần tra cứu kết quả có quyền, không tự gửi lại lệnh ghi.')
para('Các sơ đồ là tài liệu mô tả sản phẩm hiện tại, không là kiểm thử mới hoặc nghiệm thu production. M0–M3 và M4 reports/CSV/audit đã hoàn thành theo checkpoint; recovery inputs và UI gate4 giữ nguyên. Không sửa code, schema, grants hoặc báo cáo nguồn trong lượt này.')
doc.add_page_break();doc.add_heading('Những thay đổi so với bảy sơ đồ cũ',1)
table(['Sơ đồ cũ','Vấn đề từ hình nguồn','Nội dung cập nhật'],changes,[3.4,10.1,13.2])
para('Đã đọc trực tiếp cả bảy hình nguồn: tổng thể, tài khoản, tuyển sinh, học phí, đào tạo, giảng dạy và báo cáo. Điểm sai quan trọng: payment không tạo Invoice; Student không thanh toán online hoặc tự xác nhận tiền; Tuyển sinh không đăng ký thay; Teacher không tự chốt bảng điểm.')
for item in diagrams:
 doc.add_page_break();doc.add_heading(item['code']+' '+item['title'],1)
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(2);p.add_run().add_picture(str(ASSETS/(item['code']+'.png')),width=Cm(22))
 para(item['scope'])
for i in range(0,len(diagrams),5):
 doc.add_page_break();doc.add_heading('Điều kiện nghiệp vụ của các sơ đồ '+str(i//5+1),1)
 table(['Sơ đồ','Điều kiện và giới hạn cần giữ'],[[item['code']+' '+item['title'],item['guards']] for item in diagrams[i:i+5]],[6.3,20.4])
for i in range(0,len(trace),13):
 doc.add_page_break();doc.add_heading('Truy vết chức năng nghiệp vụ '+str(i//13+1),1)
 table(['Mã NV','Chức năng đã chấp nhận','Sơ đồ cập nhật','Màn hình','Workflow'],[['NV'+a,b,c,e,f] for a,b,c,e,f in trace[i:i+13]],[2,10,5,4,5.7])
 para('NV10/NV11 và F15 đã loại bỏ, không đánh lại số. Có 37 chức năng NV còn hiệu lực, 32 mã F và 21 W; số use case/oval không bắt buộc bằng số NV, màn hình hoặc endpoint. Thao tác chung logout/password/profile được biểu diễn tại UC01; Student home là cách trình bày dữ liệu đã có.')
doc.add_page_break();doc.add_heading('Nguồn đối chiếu và lưu ý mô hình hóa',1)
table(['Nội dung','Nguồn của dự án'],[
 ['Quyền và quy tắc chuẩn','docs/requirements.md; docs/design/FUNCTION_SPEC.md; docs/DECISIONS.md'],
 ['Danh mục và workflow','docs/design/FUNCTION_CATALOG.md; docs/design/WORKFLOW_SPEC.md'],
 ['Hợp đồng và dữ liệu','docs/design/E04_DTO_SPEC.md; E04_FINAL_CONTRACT.md; E04_REPLAY_SPEC.md'],
 ['Trạng thái hiện tại','AGENTS.md; docs/CURRENT_STATE.md; M4_CLOSURE_MATRIX.md'],
 ['Triển khai chức năng','LocalRuntimeServices.cs và các service/endpoint Enrollment, Academics, Classes, Admissions, Finance, Reports, Administration đã được đối chiếu trong các lượt phân tích nghiệp vụ.'],
 ['Nguồn sơ đồ cũ','docs/Nhom18_ThietKeXuLy.docx, phần V, bảy hình use case. SHA256: 40B87ABBDF8C5CF65541827CA21B4240E7D3EB8985589505B7DA6D4C57DF2B4B.'],
],[6,20.7])
para('NV39 không được vẽ thành một thao tác người dùng “ghi log”. Mandatory audit đi cùng các lệnh bắt buộc theo danh mục; security-only events best-effort là nhánh riêng. Không include audit bắt buộc vào mọi read hoặc đồng nhất login/logout logging với mutation transaction.')
para('UC07 thể hiện include tính tổng chính thức từ finalization. Teacher vẫn xem tổng nháp từ điểm đã lưu, thiếu điểm hiện “Chưa đủ điểm”; đó là phép tính khi đọc, không tự lưu final_score hoặc công bố. UC06 không vẽ Class COMPLETED dẫn thẳng đến Student result; kết quả chỉ mở sau UC07 FINALIZED.')
para('Các hình giữ mô hình mục tiêu người dùng; cookie/CSRF/SQL locks/rowversion là cơ chế hỗ trợ, không phải tác nhân hoặc use case nghiệp vụ mới. Source SVG và PlantUML của mỗi hình được bàn giao để chỉnh sửa độc lập; bố cục PlantUML có thể khác bản SVG/PDF nhưng các actor, association và include cùng một mô hình.')
out=ROOT/'docs/USE_CASE_CAP_NHAT_THEO_DU_AN_20261014.docx';doc.save(out)
(DIR/'trace.json').write_text(json.dumps(trace,ensure_ascii=False,indent=2),encoding='utf-8')
(DIR/'README.md').write_text('# Bộ sơ đồ use case cập nhật\n\nUC00 tổng quan; UC01–UC09 chi tiết. Mỗi hình có SVG vector, PNG, PDF riêng và nguồn PlantUML UTF-8. model.json lưu actor/association/include; trace.json truy vết 37 NV còn hiệu lực. Không có NV10/NV11, không tác nhân Hệ thống.\n\nBản Word và PDF tổng hợp nằm tại docs/USE_CASE_CAP_NHAT_THEO_DU_AN_20261014. Không sửa báo cáo nguồn.\n',encoding='utf-8')
with zipfile.ZipFile(ROOT/'docs/USE_CASE_NGUON_SO_DO_20261014.zip','w',zipfile.ZIP_DEFLATED) as z:
 for path in sorted(ASSETS.iterdir()):z.write(path,'diagrams/'+path.name)
 for name in ['model.json','trace.json','README.md']:z.write(DIR/name,name)
print(out)
print('Diagrams',len(diagrams),'Trace',len(trace))
